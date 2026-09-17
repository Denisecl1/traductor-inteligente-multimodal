import io
import json
import os

from email.parser import BytesParser
from email.policy import default
from http.server import BaseHTTPRequestHandler

from docx import Document
from pypdf import PdfReader
from openai import OpenAI


# =========================================
# CONFIGURACIÓN
# =========================================

MODEL = os.getenv(
    "OPENAI_MODEL",
    "gpt-5.6-luna"
)


# Límite del archivo original
MAX_DOCUMENT_BYTES = 4 * 1024 * 1024

# Límite aproximado de la petición
# completa enviada a Vercel.
MAX_REQUEST_BYTES = 4_400_000

# Evitamos documentos excesivamente largos
# para esta versión del proyecto.
MAX_EXTRACTED_CHARACTERS = 18000

# Dividimos documentos largos en fragmentos.
TRANSLATION_CHUNK_SIZE = 4000


ALLOWED_EXTENSIONS = {
    "pdf",
    "docx",
    "txt"
}


ALLOWED_MIME_TYPES = {
    "application/pdf",

    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",

    "text/plain",

    # Algunos navegadores pueden enviar
    # este tipo genérico.
    "application/octet-stream"
}


# =========================================
# EXCEPCIONES PERSONALIZADAS
# =========================================

class DocumentContentError(Exception):
    """
    Error cuando el documento existe,
    pero no contiene texto utilizable.
    """

    pass


class UnsupportedLanguageError(Exception):
    """
    Error cuando el texto no está
    principalmente en español o inglés.
    """

    pass


# =========================================
# VALIDACIÓN DEL DOCUMENTO
# =========================================

class DocumentValidator:

    @staticmethod
    def get_extension(filename):

        if not isinstance(filename, str):

            return ""


        parts = (
            filename
            .lower()
            .rsplit(".", 1)
        )


        if len(parts) != 2:

            return ""


        return parts[1]


    @staticmethod
    def validate(
        filename,
        mime_type,
        document_bytes
    ):

        # ---------------------------------
        # NOMBRE
        # ---------------------------------

        if not filename:

            return (
                False,
                "El documento no tiene un nombre válido."
            )


        extension = (
            DocumentValidator
            .get_extension(
                filename
            )
        )


        # ---------------------------------
        # EXTENSIÓN
        # ---------------------------------

        if extension not in ALLOWED_EXTENSIONS:

            return (
                False,
                "El formato del documento no está permitido."
            )


        # ---------------------------------
        # MIME TYPE
        # ---------------------------------

        if (
            mime_type
            and
            mime_type not in ALLOWED_MIME_TYPES
        ):

            return (
                False,
                "El tipo de archivo del documento no está permitido."
            )


        # ---------------------------------
        # CONTENIDO
        # ---------------------------------

        if not isinstance(
            document_bytes,
            bytes
        ):

            return (
                False,
                "El documento recibido no es válido."
            )


        if len(document_bytes) == 0:

            return (
                False,
                "El documento seleccionado está vacío."
            )


        # ---------------------------------
        # TAMAÑO
        # ---------------------------------

        if (
            len(document_bytes)
            >
            MAX_DOCUMENT_BYTES
        ):

            return (
                False,
                "El documento supera el límite de 4 MB."
            )


        return (
            True,
            extension
        )


# =========================================
# PARSER MULTIPART
# =========================================

class MultipartDocumentParser:

    @staticmethod
    def parse(
        content_type,
        raw_body
    ):

        mime_message = (

            (
                f"Content-Type: {content_type}\r\n"
                f"MIME-Version: 1.0\r\n"
                f"\r\n"
            )
            .encode("utf-8")

            +

            raw_body
        )


        try:

            message = (
                BytesParser(
                    policy=default
                )
                .parsebytes(
                    mime_message
                )
            )

        except Exception:

            raise ValueError(
                "No fue posible interpretar el documento enviado."
            )


        if not message.is_multipart():

            raise ValueError(
                "La solicitud multipart no es válida."
            )


        for part in message.iter_parts():

            disposition = (
                part
                .get_content_disposition()
            )


            if disposition != "form-data":

                continue


            field_name = (
                part.get_param(
                    "name",
                    header="content-disposition"
                )
            )


            if field_name != "document":

                continue


            filename = (
                part.get_filename()
                or
                "documento"
            )


            mime_type = (
                part.get_content_type()
            )


            document_bytes = (
                part.get_payload(
                    decode=True
                )
                or
                b""
            )


            return (
                filename,
                mime_type,
                document_bytes
            )


        raise ValueError(
            "No se encontró el documento en la solicitud."
        )


# =========================================
# EXTRACTOR DE TEXTO
# =========================================

class DocumentTextExtractor:

    @staticmethod
    def extract(
        extension,
        document_bytes
    ):

        if extension == "txt":

            text = (
                DocumentTextExtractor
                .extract_txt(
                    document_bytes
                )
            )


        elif extension == "pdf":

            text = (
                DocumentTextExtractor
                .extract_pdf(
                    document_bytes
                )
            )


        elif extension == "docx":

            text = (
                DocumentTextExtractor
                .extract_docx(
                    document_bytes
                )
            )


        else:

            raise DocumentContentError(
                "El formato del documento no puede procesarse."
            )


        text = (
            DocumentTextExtractor
            .clean_text(
                text
            )
        )


        if not text:

            raise DocumentContentError(
                "El documento no contiene texto utilizable."
            )


        if (
            len(text)
            >
            MAX_EXTRACTED_CHARACTERS
        ):

            raise DocumentContentError(
                "El documento contiene demasiado texto para procesarlo en esta versión. Usa un archivo más corto."
            )


        return text


    # =====================================
    # TXT
    # =====================================

    @staticmethod
    def extract_txt(
        document_bytes
    ):

        encodings = [
            "utf-8",
            "utf-8-sig",
            "latin-1"
        ]


        for encoding in encodings:

            try:

                return (
                    document_bytes
                    .decode(
                        encoding
                    )
                )

            except UnicodeDecodeError:

                continue


        raise DocumentContentError(
            "No fue posible leer el archivo TXT."
        )


    # =====================================
    # PDF
    # =====================================

    @staticmethod
    def extract_pdf(
        document_bytes
    ):

        try:

            reader = PdfReader(
                io.BytesIO(
                    document_bytes
                )
            )


            pages = []


            for page_number, page in enumerate(
                reader.pages,
                start=1
            ):

                page_text = (
                    page.extract_text()
                    or
                    ""
                )


                page_text = (
                    page_text
                    .strip()
                )


                if page_text:

                    pages.append(
                        page_text
                    )


            return "\n\n".join(
                pages
            )


        except Exception:

            raise DocumentContentError(
                "No fue posible leer el archivo PDF."
            )


    # =====================================
    # WORD DOCX
    # =====================================

    @staticmethod
    def extract_docx(
        document_bytes
    ):

        try:

            document = Document(
                io.BytesIO(
                    document_bytes
                )
            )


            blocks = []


            # -----------------------------
            # PÁRRAFOS Y TÍTULOS
            # -----------------------------

            for paragraph in document.paragraphs:

                text = (
                    paragraph.text
                    .strip()
                )


                if not text:

                    continue


                style_name = ""

                try:

                    style_name = (
                        paragraph.style.name
                        or
                        ""
                    )

                except Exception:

                    style_name = ""


                # Conservamos el orden.
                # No modificamos el contenido
                # del encabezado.

                if (
                    style_name
                    .lower()
                    .startswith(
                        "heading"
                    )
                    or
                    style_name
                    .lower()
                    ==
                    "title"
                ):

                    blocks.append(
                        text
                    )

                else:

                    blocks.append(
                        text
                    )


            # -----------------------------
            # TABLAS
            # -----------------------------

            for table in document.tables:

                for row in table.rows:

                    cells = []


                    for cell in row.cells:

                        cell_text = (
                            cell.text
                            .strip()
                        )


                        if cell_text:

                            cells.append(
                                cell_text
                            )


                    if cells:

                        blocks.append(
                            " | ".join(
                                cells
                            )
                        )


            return "\n\n".join(
                blocks
            )


        except Exception:

            raise DocumentContentError(
                "No fue posible leer el archivo Word."
            )


    # =====================================
    # LIMPIAR TEXTO
    # =====================================

    @staticmethod
    def clean_text(
        text
    ):

        if not isinstance(text, str):

            return ""


        lines = (
            text
            .replace(
                "\r\n",
                "\n"
            )
            .replace(
                "\r",
                "\n"
            )
            .split(
                "\n"
            )
        )


        cleaned_lines = []


        for line in lines:

            cleaned_lines.append(
                line.rstrip()
            )


        text = "\n".join(
            cleaned_lines
        )


        # Evitar demasiados saltos
        # consecutivos.

        while "\n\n\n" in text:

            text = text.replace(
                "\n\n\n",
                "\n\n"
            )


        return text.strip()


# =========================================
# SERVICIO DE TRADUCCIÓN
# =========================================

class DocumentTranslationService:

    def __init__(self):

        api_key = os.getenv(
            "OPENAI_API_KEY"
        )


        if not api_key:

            raise RuntimeError(
                "OPENAI_API_KEY no está configurada."
            )


        self.client = OpenAI(
            api_key=api_key
        )


    # =====================================
    # PROCESO COMPLETO
    # =====================================

    def translate_document(
        self,
        original_text
    ):

        # ---------------------------------
        # DETECTAR IDIOMA
        # ---------------------------------

        source_language = (
            self.detect_language(
                original_text
            )
        )


        target_language = (
            "en"
            if source_language == "es"
            else "es"
        )


        # ---------------------------------
        # DIVIDIR DOCUMENTO
        # ---------------------------------

        chunks = (
            self.create_chunks(
                original_text
            )
        )


        translations = []


        # ---------------------------------
        # TRADUCIR CADA FRAGMENTO
        # ---------------------------------

        for chunk in chunks:

            translated_chunk = (
                self.translate_chunk(
                    chunk,
                    source_language
                )
            )


            translations.append(
                translated_chunk
            )


        translated_text = (
            "\n\n"
            .join(
                translations
            )
            .strip()
        )


        if not translated_text:

            raise ValueError(
                "La traducción del documento está vacía."
            )


        return {

            "source_language":
                source_language,

            "target_language":
                target_language,

            "original_text":
                original_text,

            "translation":
                translated_text
        }


    # =====================================
    # DETECTAR IDIOMA
    # =====================================

    def detect_language(
        self,
        text
    ):

        # No necesitamos enviar el documento
        # completo para detectar el idioma.

        sample = (
            text[:3000]
        )


        instructions = """
Determina el idioma principal del texto recibido.

Solamente se admiten español e inglés.

Devuelve SOLAMENTE uno de estos valores:

es
en
unsupported

Usa:
- es si el texto está principalmente en español.
- en si está principalmente en inglés.
- unsupported si está principalmente en otro idioma.

No agregues explicaciones.
"""


        response = (
            self.client
            .responses
            .create(

                model=
                    MODEL,

                instructions=
                    instructions,

                input=
                    sample,

                max_output_tokens=
                    20,

                store=False
            )
        )


        language = (
            response
            .output_text
            .strip()
            .lower()
        )


        if language == "unsupported":

            raise UnsupportedLanguageError(
                "El documento debe estar principalmente en español o inglés."
            )


        if language not in [
            "es",
            "en"
        ]:

            raise ValueError(
                "No fue posible determinar el idioma del documento."
            )


        return language


    # =====================================
    # CREAR FRAGMENTOS
    # =====================================

    def create_chunks(
        self,
        text
    ):

        paragraphs = (
            text.split(
                "\n\n"
            )
        )


        chunks = []

        current_chunk = ""


        for paragraph in paragraphs:

            paragraph = (
                paragraph.strip()
            )


            if not paragraph:

                continue


            # Si un solo párrafo es demasiado
            # grande, lo dividimos.

            if (
                len(paragraph)
                >
                TRANSLATION_CHUNK_SIZE
            ):

                if current_chunk:

                    chunks.append(
                        current_chunk.strip()
                    )

                    current_chunk = ""


                start = 0


                while (
                    start
                    <
                    len(paragraph)
                ):

                    piece = (
                        paragraph[
                            start:
                            start
                            +
                            TRANSLATION_CHUNK_SIZE
                        ]
                    )


                    chunks.append(
                        piece.strip()
                    )


                    start += (
                        TRANSLATION_CHUNK_SIZE
                    )


                continue


            candidate = (

                current_chunk
                +
                (
                    "\n\n"
                    if current_chunk
                    else ""
                )
                +
                paragraph
            )


            if (
                len(candidate)
                <=
                TRANSLATION_CHUNK_SIZE
            ):

                current_chunk = (
                    candidate
                )

            else:

                if current_chunk:

                    chunks.append(
                        current_chunk.strip()
                    )


                current_chunk = (
                    paragraph
                )


        if current_chunk:

            chunks.append(
                current_chunk.strip()
            )


        return chunks


    # =====================================
    # TRADUCIR FRAGMENTO
    # =====================================

    def translate_chunk(
        self,
        text,
        source_language
    ):

        if source_language == "es":

            target_name = "inglés"

        else:

            target_name = "español"


        instructions = f"""
Eres un traductor profesional especializado
en documentos español-inglés.

Traduce el contenido recibido al {target_name}.

REGLAS:

1. Traduce todo el contenido proporcionado.

2. Conserva el orden del texto.

3. Conserva títulos, párrafos y separaciones.

4. No resumas.

5. No elimines información.

6. No agregues información que no exista
   en el documento.

7. Mantén correctamente:
   - nombres propios
   - números
   - fechas
   - cantidades
   - precios
   - unidades
   - términos técnicos
   - siglas

8. Produce una traducción natural y comprensible.

9. No traduzcas nombres propios cuando no
   corresponda.

10. Conserva los saltos de párrafo existentes.

Devuelve SOLAMENTE el texto traducido.

No utilices Markdown adicional.
No agregues comentarios ni explicaciones.
"""


        response = (
            self.client
            .responses
            .create(

                model=
                    MODEL,

                instructions=
                    instructions,

                input=
                    text,

                max_output_tokens=
                    3000,

                store=False
            )
        )


        translated_text = (
            response
            .output_text
            .strip()
        )


        if not translated_text:

            raise ValueError(
                "La IA devolvió una traducción vacía."
            )


        return translated_text


# =========================================
# ENDPOINT DE VERCEL
# =========================================

class handler(
    BaseHTTPRequestHandler
):

    # =====================================
    # VALIDAR ORIGEN
    # =====================================

    def _origin_allowed(self):

        allowed_origin = os.getenv(
            "ALLOWED_ORIGIN",
            ""
        ).rstrip("/")


        request_origin = (
            self.headers
            .get(
                "Origin",
                ""
            )
            .rstrip("/")
        )


        if not allowed_origin:

            return False


        return (
            request_origin
            ==
            allowed_origin
        )


    # =====================================
    # CORS
    # =====================================

    def _add_cors_headers(self):

        origin = (
            self.headers
            .get(
                "Origin",
                ""
            )
        )


        if self._origin_allowed():

            self.send_header(
                "Access-Control-Allow-Origin",
                origin
            )


            self.send_header(
                "Vary",
                "Origin"
            )


    # =====================================
    # RESPUESTA JSON
    # =====================================

    def _send_json(
        self,
        status_code,
        data
    ):

        body = (
            json.dumps(
                data,
                ensure_ascii=False
            )
            .encode(
                "utf-8"
            )
        )


        self.send_response(
            status_code
        )


        self._add_cors_headers()


        self.send_header(
            "Content-Type",
            "application/json; charset=utf-8"
        )


        self.send_header(
            "Content-Length",
            str(
                len(body)
            )
        )


        self.end_headers()


        self.wfile.write(
            body
        )


    # =====================================
    # LEER BODY
    # =====================================

    def _read_body(self):

        try:

            content_length = int(
                self.headers.get(
                    "Content-Length",
                    "0"
                )
            )

        except ValueError:

            raise ValueError(
                "El tamaño de la solicitud no es válido."
            )


        if content_length <= 0:

            raise ValueError(
                "La solicitud está vacía."
            )


        if (
            content_length
            >
            MAX_REQUEST_BYTES
        ):

            raise OverflowError(
                "El documento supera el tamaño permitido."
            )


        return (
            self.rfile.read(
                content_length
            )
        )


    # =====================================
    # OPTIONS
    # =====================================

    def do_OPTIONS(self):

        if not self._origin_allowed():

            self._send_json(
                403,
                {
                    "error":
                        "Origen no autorizado."
                }
            )

            return


        self.send_response(
            204
        )


        self._add_cors_headers()


        self.send_header(
            "Access-Control-Allow-Methods",
            "POST, OPTIONS"
        )


        self.send_header(
            "Access-Control-Allow-Headers",
            "Content-Type"
        )


        self.send_header(
            "Access-Control-Max-Age",
            "86400"
        )


        self.end_headers()


    # =====================================
    # POST /api/documents
    # =====================================

    def do_POST(self):

        # ---------------------------------
        # ORIGEN
        # ---------------------------------

        if not self._origin_allowed():

            self._send_json(
                403,
                {
                    "error":
                        "Origen no autorizado."
                }
            )

            return


        content_type_header = (
            self.headers
            .get(
                "Content-Type",
                ""
            )
        )


        media_type = (
            content_type_header
            .split(";")[0]
            .strip()
            .lower()
        )


        # ---------------------------------
        # SOLO MULTIPART
        # ---------------------------------

        if (
            media_type
            !=
            "multipart/form-data"
        ):

            self._send_json(
                415,
                {
                    "error":
                        (
                            "El documento debe enviarse "
                            "como multipart/form-data."
                        )
                }
            )

            return


        try:

            # -----------------------------
            # LEER PETICIÓN
            # -----------------------------

            raw_body = (
                self._read_body()
            )


            # -----------------------------
            # EXTRAER DOCUMENTO
            # -----------------------------

            (
                filename,
                mime_type,
                document_bytes

            ) = (
                MultipartDocumentParser
                .parse(
                    content_type_header,
                    raw_body
                )
            )


            # -----------------------------
            # VALIDAR DOCUMENTO
            # -----------------------------

            (
                is_valid,
                validation_result

            ) = (
                DocumentValidator
                .validate(
                    filename,
                    mime_type,
                    document_bytes
                )
            )


            if not is_valid:

                self._send_json(
                    400,
                    {
                        "error":
                            validation_result
                    }
                )

                return


            extension = (
                validation_result
            )


            # -----------------------------
            # EXTRAER TEXTO
            # -----------------------------

            original_text = (
                DocumentTextExtractor
                .extract(
                    extension,
                    document_bytes
                )
            )


            # -----------------------------
            # TRADUCIR
            # -----------------------------

            service = (
                DocumentTranslationService()
            )


            result = (
                service
                .translate_document(
                    original_text
                )
            )


            # -----------------------------
            # RESPUESTA
            # -----------------------------

            self._send_json(
                200,
                {
                    "filename":
                        filename,

                    "original_text":
                        result[
                            "original_text"
                        ],

                    "translation":
                        result[
                            "translation"
                        ],

                    "source_language":
                        result[
                            "source_language"
                        ],

                    "target_language":
                        result[
                            "target_language"
                        ]
                }
            )


        # ---------------------------------
        # ARCHIVO DEMASIADO GRANDE
        # ---------------------------------

        except OverflowError as error:

            self._send_json(
                413,
                {
                    "error":
                        str(error)
                }
            )


        # ---------------------------------
        # DOCUMENTO VACÍO / SIN TEXTO
        # ---------------------------------

        except DocumentContentError as error:

            self._send_json(
                422,
                {
                    "error":
                        str(error)
                }
            )


        # ---------------------------------
        # IDIOMA NO ADMITIDO
        # ---------------------------------

        except UnsupportedLanguageError as error:

            self._send_json(
                422,
                {
                    "error":
                        str(error)
                }
            )


        # ---------------------------------
        # SOLICITUD / RESPUESTA INVÁLIDA
        # ---------------------------------

        except ValueError as error:

            print(
                "Error procesando documento:",
                type(error).__name__
            )


            self._send_json(
                400,
                {
                    "error":
                        str(error)
                }
            )


        # ---------------------------------
        # ERROR INTERNO
        # ---------------------------------

        except Exception as error:

            print(
                "Error interno en documentos:",
                type(error).__name__
            )


            self._send_json(
                500,
                {
                    "error":
                        (
                            "No fue posible procesar "
                            "el documento en este momento."
                        )
                }
            )


    # =====================================
    # GET NO PERMITIDO
    # =====================================

    def do_GET(self):

        self._send_json(
            405,
            {
                "error":
                    "Método no permitido."
            }
        )